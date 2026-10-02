import json
import unittest
from pathlib import Path

from docx import Document

from project_status_report.render_docx import render_report
from project_status_report.rules import calculate_progress


ROOT = Path(__file__).resolve().parents[1]


class OperationalItemsTests(unittest.TestCase):
    def setUp(self):
        self.data = json.loads(
            (ROOT / "examples" / "sample-report-data.json").read_text(encoding="utf-8")
        )

    def test_operational_items_render_but_never_affect_progress(self):
        baseline = calculate_progress(self.data)
        self.assertEqual(baseline.earned, 72000)
        self.assertEqual(baseline.basis, "stages")

        self.data["operational_items"][0]["budget"] = 50000000
        changed = calculate_progress(self.data)
        self.assertEqual(changed.earned, baseline.earned)
        self.assertEqual(changed.items, baseline.items)

        temp_root = ROOT / "output" / "test-artifacts"
        temp_root.mkdir(parents=True, exist_ok=True)
        path = temp_root / "operational-clean.docx"
        render_report(self.data, path, include_comments=False)
        text = "\n".join(
            paragraph.text for table in Document(path).tables
            for row in table.rows for cell in row.cells
            for paragraph in cell.paragraphs
        )
        self.assertIn("Оперативная задача", text)
        self.assertIn(self.data["operational_items"][0]["title"], text)


if __name__ == "__main__":
    unittest.main()
