import json
import unittest
from pathlib import Path

from docx import Document

from project_status_report.render_docx import render_report

ROOT = Path(__file__).resolve().parents[1]


class ReportLayoutTests(unittest.TestCase):
    def test_v060_first_page_layout(self):
        data = json.loads((ROOT / "examples" / "sample-report-data.json").read_text(encoding="utf-8"))
        data["report"]["project_start_date"] = "2026-08-15"
        data["report"]["additional_agreements"] = [
            "ДС № 1 от 01.09.2026",
            "ДС № 2 от 15.09.2026",
        ]
        data["stages"][0]["actual_start"] = "2026-09-02"
        data["stages"][0]["start_deviation"] = "+1 кал. день"
        temp_root = ROOT / "output" / "test-artifacts"
        temp_root.mkdir(parents=True, exist_ok=True)
        path = temp_root / "layout-clean.docx"
        render_report(data, path, include_comments=False)
        document = Document(path)

        paragraphs = [p.text for p in document.paragraphs]
        headings = {p.text: p.style.name for p in document.paragraphs if p.text}
        table_text = "\n".join(
            p.text for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs
        )

        self.assertEqual(headings["1. Общие сведения"], "Heading 1")
        self.assertEqual(headings["Сводные данные по проекту"], "Heading 2")
        self.assertEqual(headings["План и статус этапов / дополнительных соглашений"], "Heading 2")
        self.assertEqual(headings["Общий статус проекта"], "Heading 2")
        self.assertEqual(headings["Контроль состояния проекта"], "Heading 2")
        self.assertNotIn("Health check", paragraphs)

        passport = document.tables[0]
        self.assertEqual(len(passport.columns), 2)
        self.assertGreaterEqual(len(passport.rows), 7)
        labels = [row.cells[0].text for row in passport.rows]
        self.assertEqual(labels[:7], [
            "Заказчик", "Проект", "Руководитель проекта от Исполнителя",
            "Руководитель проекта от Заказчика", "Основание", "Номер отчета", "Отчетный период"
        ])
        self.assertIn("Дата начала проекта", table_text)
        self.assertIn("2026-08-15", table_text)
        self.assertIn("ДС № 2 от 15.09.2026", table_text)
        self.assertIn("План начала", table_text)
        self.assertIn("Факт начала", table_text)
        self.assertIn("Освоено / %", table_text)


if __name__ == "__main__":
    unittest.main()
