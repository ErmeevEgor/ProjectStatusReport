from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .rules import (
    calculate_progress,
    coefficient,
    inherit_parent_periods,
    page2_contractual_rows,
    progress_formula,
    resolve_payment_display,
)


def _border(cell, color: str = "B7B7B7", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _margins(cell, twips: int = 20) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")


def _fill(cell, color: str = "E7E6E6") -> None:
    """Apply cosmetic table shading without using it as a status signal."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)


def _table(doc: Document, rows: int, cols: int, widths_mm: list[float]):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            _border(cell)
            _margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index < len(widths_mm):
                cell.width = Mm(widths_mm[index])
    return table


def _comment_text(
    source_ids: list[str] | None, source_map: dict[str, dict[str, Any]]
) -> str | None:
    if not source_ids:
        return None
    parts: list[str] = []
    for source_id in source_ids:
        source = source_map.get(source_id)
        if not source:
            continue
        lines = [f"Источник: {source.get('name', source_id)}"]
        if source.get("date"):
            lines.append(f"Дата: {source['date']}")
        if source.get("reference"):
            lines.append(f"Ссылка/место: {source['reference']}")
        lines.append(f"Подтверждение: {source.get('confidence', 'unknown')}")
        if source.get("note"):
            lines.append(str(source["note"]))
        parts.append("\n".join(lines))
    return "\n\n".join(parts) if parts else None


def _cell(
    doc: Document,
    cell,
    text: Any,
    *,
    bold: bool = False,
    size: float = 6.4,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    source_ids: list[str] | None = None,
    source_map: dict[str, dict[str, Any]] | None = None,
    comments: bool = False,
    extra_comment: str | None = None,
    line_spacing: float = 1.0,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "Normal"
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = line_spacing
    run = paragraph.add_run("—" if text in (None, "") else str(text))
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    if comments:
        comment = _comment_text(source_ids, source_map or {})
        if extra_comment:
            comment = f"{comment}\n\n{extra_comment}" if comment else extra_comment
        if comment:
            doc.add_comment(run, text=comment, author="OSP Skill", initials="OSP")


def _header_row(doc: Document, row, labels: list[str], size: float = 6.0) -> None:
    for index, label in enumerate(labels):
        _fill(row.cells[index])
        _cell(
            doc,
            row.cells[index],
            label,
            bold=True,
            size=size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )


def _heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True


def _paragraph(
    doc: Document,
    text: Any,
    *,
    bold_lead: str | None = None,
    source_ids: list[str] | None = None,
    source_map: dict[str, dict[str, Any]] | None = None,
    comments: bool = False,
) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(2)
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        lead.font.color.rgb = RGBColor(0, 0, 0)
    run = paragraph.add_run("—" if text in (None, "") else str(text))
    run.font.color.rgb = RGBColor(0, 0, 0)
    if comments:
        comment = _comment_text(source_ids, source_map or {})
        if comment:
            doc.add_comment(run, text=comment, author="OSP Skill", initials="OSP")


def _bold_label(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)


def _setup(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(7)
    section.bottom_margin = Mm(7)
    section.left_margin = Mm(7)
    section.right_margin = Mm(7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(7)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    for name, size, before, after in (
        ("Heading 1", 13, 0, 3),
        ("Heading 2", 9, 3, 2),
        ("Heading 3", 8, 2, 1),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(7)
    bullet.font.color.rgb = RGBColor(0, 0, 0)
    bullet.paragraph_format.space_after = Pt(0)


def _money(value: Any) -> str:
    if value in (None, ""):
        return "—"
    return f"{float(value):,.0f}".replace(",", " ")


def _percent(value: float) -> str:
    return str(f"{value:.1f}").replace(".", ",") + "%"


def render_report(
    data: dict[str, Any], out_path: str | Path, *, include_comments: bool = True
) -> Path:
    out_path = Path(out_path)
    data = inherit_parent_periods(data)
    doc = Document()
    _setup(doc)
    source_map = {source["id"]: source for source in data.get("sources", []) if source.get("id")}
    report = data["report"]
    field_sources = report.get("field_sources") or {}
    progress = calculate_progress(data)

    _heading(doc, "1. Общие сведения", 1)
    passport = _table(doc, 4, 4, [54, 84, 54, 91])
    passport_rows = [
        ("Заказчик", "customer", "Проект", "project_name"),
        ("Руководитель проекта от Заказчика", "customer_pm", "Основание", "contract_basis"),
        ("Дополнительное соглашение", "additional_agreement", "Номер отчета", "report_number"),
        ("Руководитель проекта от Исполнителя", "contractor_pm", "Отчетный период", "reporting_period"),
    ]
    for row_index, (label_1, key_1, label_2, key_2) in enumerate(passport_rows):
        _fill(passport.cell(row_index, 0))
        _fill(passport.cell(row_index, 2))
        _cell(doc, passport.cell(row_index, 0), label_1, bold=True, size=6.2)
        value_1 = report.get(key_1)
        if key_1 == "report_number" and value_1 not in (None, ""):
            value_1 = f"№ {value_1}"
        _cell(doc, passport.cell(row_index, 1), value_1, size=6.2, source_ids=field_sources.get(key_1), source_map=source_map, comments=include_comments)
        _cell(doc, passport.cell(row_index, 2), label_2, bold=True, size=6.2)
        value_2 = report.get(key_2)
        if key_2 == "report_number" and value_2 not in (None, ""):
            value_2 = f"№ {value_2}"
        _cell(doc, passport.cell(row_index, 3), value_2, size=6.2, source_ids=field_sources.get(key_2), source_map=source_map, comments=include_comments)

    _heading(doc, "Сводные данные и общий статус", 2)
    summary = _table(doc, 2, 4, [54, 84, 54, 91])
    summary_rows = [
        ("Плановый период этапа", report.get("planned_stage_period"), "Факт", report.get("actual_stage_period")),
        ("Этап проекта", report.get("stage_name"), "Расчетный прогресс", _percent(progress.percent)),
    ]
    for row_index, values in enumerate(summary_rows):
        for column, value in enumerate(values):
            if column in {0, 2}:
                _fill(summary.cell(row_index, column))
            _cell(doc, summary.cell(row_index, column), value, bold=column in {0, 2}, size=6.3)

    _heading(doc, "Общий статус проекта", 2)
    for point in report.get("overall_status_points", []):
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(point))
        run.font.color.rgb = RGBColor(0, 0, 0)

    _heading(doc, "План и статус этапов", 2)
    stages = data.get("stages", [])
    stage_table = _table(doc, 1 + len(stages), 7, [63, 24, 28, 29, 20, 31, 88])
    _header_row(doc, stage_table.rows[0], ["Этап / подэтап", "Трудоемкость", "Бюджет", "Статус", "Коэф.", "Освоение", "Комментарий"], 5.8)
    for row_index, item in enumerate(stages, 1):
        budget = item.get("budget")
        # Stage rows are descriptive when progress is calculated from the
        # complete task-level baseline. Showing a second coefficient here would
        # imply double counting and can contradict the deterministic total.
        factor = (
            coefficient(item.get("status"))
            if budget is not None and progress.basis == "stages"
            else None
        )
        earned = float(budget) * factor if budget is not None and factor is not None else None
        values = [item.get("title"), item.get("duration"), _money(budget), item.get("status"), f"{factor * 100:.0f}%" if factor is not None else "—", _money(earned), item.get("comment")]
        for column, value in enumerate(values):
            _cell(doc, stage_table.cell(row_index, column), value, size=5.65, source_ids=item.get("sources"), source_map=source_map, comments=include_comments)

    _paragraph(doc, f"бюджет {_money(progress.approved_budget)} ₽ | расчетное освоение {_money(progress.earned)} ₽ | остаток {_money(progress.approved_budget - progress.earned)} ₽ | {_percent(progress.percent)}", bold_lead="Итоговый прогресс: ")

    _heading(doc, "План / факт оплат", 2)
    payments = data.get("payments", [])
    payment_table = _table(doc, 1 + len(payments), 7, [43, 27, 68, 38, 42, 29, 36])
    _header_row(doc, payment_table.rows[0], ["Платеж", "Сумма", "План / основание по ДС", "Плановая дата / срок", "Факт", "Дата факта", "Статус"], 5.45)
    for row_index, payment in enumerate(payments, 1):
        display = resolve_payment_display(payment, report.get("report_date"))
        values = [payment.get("name"), _money(payment.get("amount")), payment.get("plan_basis"), payment.get("planned_date"), display.actual_text, display.actual_date, display.status]
        for column, value in enumerate(values):
            _cell(doc, payment_table.cell(row_index, column), value, size=5.35, source_ids=payment.get("sources"), source_map=source_map, comments=include_comments)

    doc.add_page_break()
    _heading(doc, "2. Данные по задачам", 1)
    agreement = report.get("additional_agreement") or report.get("contract_basis") or "договору"
    _paragraph(doc, f"Страница отражает договорный состав работ по {agreement}. Операционные задачи не участвуют в расчете прогресса.")

    tasks = page2_contractual_rows(data)
    task_table = _table(doc, 1 + len(tasks), 10, [10, 56, 24, 24, 23, 28, 24, 25, 25, 44])
    _header_row(doc, task_table.rows[0], ["№", "Задача / результат", "Ответственный", "Статус", "Бюджет", "Расчетное освоение", "Плановое начало", "Плановое завершение", "Фактическое завершение", "Комментарий"], 4.7)
    for row_index, item in enumerate(tasks, 1):
        budget = item.get("budget")
        earned = float(budget) * coefficient(item.get("status")) if budget is not None else None
        task_text = item.get("title")
        if item.get("result"):
            task_text = f"{task_text} → {item.get('result')}"
        values = [item.get("id"), task_text, item.get("owner"), item.get("status"), _money(budget), _money(earned), item.get("planned_start"), item.get("planned_end"), item.get("actual_end"), item.get("comment")]
        for column, value in enumerate(values):
            inherited_note = None
            if column in {6, 7} and item.get("date_basis") == "parent_stage_period":
                inherited_note = f"Плановый период унаследован от договорного родительского этапа {item.get('contract_parent_id') or item.get('parent_id')}."
            _cell(doc, task_table.cell(row_index, column), value, size=4.4, source_ids=item.get("sources"), source_map=source_map, comments=include_comments, extra_comment=inherited_note, line_spacing=0.84)

    _heading(doc, "Оперативные факты отчетного периода (не входят в расчет прогресса)", 2)
    operational = data.get("operational_items", [])
    operational_table = _table(doc, 1 + len(operational), 6, [58, 65, 38, 33, 37, 52])
    _header_row(doc, operational_table.rows[0], ["Оперативная задача", "Результат", "Ответственный", "Статус", "Срок / ориентир", "Комментарий"], 5.0)
    for row_index, item in enumerate(operational, 1):
        period = item.get("planned_end") or item.get("duration")
        if item.get("planned_start") and item.get("planned_end"):
            period = f"{item.get('planned_start')}–{item.get('planned_end')}"
        values = [item.get("title"), item.get("result"), item.get("owner"), item.get("status"), period, item.get("comment")]
        for column, value in enumerate(values):
            _cell(doc, operational_table.cell(row_index, column), value, size=4.7, source_ids=item.get("sources"), source_map=source_map, comments=include_comments, line_spacing=0.9)

    _paragraph(doc, progress_formula(data))

    doc.add_page_break()
    _heading(doc, "3. Риски проекта", 1)
    risks = data.get("risks", [])
    risk_table = _table(doc, 1 + len(risks), 9, [19, 18, 40, 23, 22, 46, 46, 45, 24])
    _header_row(doc, risk_table.rows[0], ["Дата", "Тип", "Отклонение / риск", "Вероятность / факт", "Влияние", "Влияние на проект", "Корректирующие мероприятия", "Результат / принятое решение", "Статус"], 5.05)
    for row_index, risk in enumerate(risks, 1):
        values = [risk.get("date"), risk.get("type"), risk.get("title"), risk.get("probability_or_fact"), risk.get("impact_degree"), risk.get("impact"), risk.get("actions"), risk.get("result"), risk.get("status")]
        for column, value in enumerate(values):
            _cell(doc, risk_table.cell(row_index, column), value, size=4.9, source_ids=risk.get("sources"), source_map=source_map, comments=include_comments)

    _bold_label(doc, "Ключевой вывод по рискам")
    _paragraph(doc, report.get("risk_summary"))
    if report.get("financial_observation") not in (None, ""):
        _bold_label(doc, "Финансовое наблюдение")
        _paragraph(doc, report.get("financial_observation"))

    doc.add_page_break()
    _heading(doc, "4. Ключевые вопросы и проблемы", 1)
    _heading(doc, "Health check", 2)
    health = data.get("health_check", [])
    health_table = _table(doc, 1 + len(health), 4, [12, 88, 31, 152])
    _header_row(doc, health_table.rows[0], ["№", "Контрольный вопрос", "Ответ", "Комментарий"], 5.9)
    for row_index, item in enumerate(health, 1):
        values = [item.get("number"), item.get("question"), item.get("answer"), item.get("comment")]
        for column, value in enumerate(values):
            _cell(doc, health_table.cell(row_index, column), value, size=5.6, source_ids=item.get("sources"), source_map=source_map, comments=include_comments)

    _heading(doc, "Открытые вопросы / действия", 2)
    open_items = data.get("open_items", [])
    open_table = _table(doc, 1 + len(open_items), 5, [69, 101, 39, 31, 43])
    _header_row(doc, open_table.rows[0], ["Открытый вопрос / действие", "Необходимое действие / решение", "Ответственный", "Срок", "Статус / результат"], 5.65)
    for row_index, item in enumerate(open_items, 1):
        values = [item.get("item"), item.get("required_action"), item.get("owner"), item.get("due_date"), item.get("status")]
        for column, value in enumerate(values):
            _cell(doc, open_table.cell(row_index, column), value, size=5.45, source_ids=item.get("sources"), source_map=source_map, comments=include_comments)

    _paragraph(doc, report.get("next_control_milestone"), bold_lead="Следующий контрольный ориентир: ")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
