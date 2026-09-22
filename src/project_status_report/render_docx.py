from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .rules import calculate_progress, coefficient, inherit_parent_periods, page2_contractual_rows

LIGHT = "F2F2F2"
HEADER = "D9EAF7"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
RED = "F4CCCC"
DARKBLUE = "1F4E78"


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def _border(cell, color: str = "B7B7B7", size: str = "4") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _margins(cell, twips: int = 32) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side in ["top", "left", "bottom", "right"]:
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")


def _compact(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            _border(cell)
            _margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _widths(table, values_mm: list[float]) -> None:
    for row in table.rows:
        for i, width in enumerate(values_mm):
            if i < len(row.cells):
                row.cells[i].width = Mm(width)


def _comment_text(source_ids: list[str] | None, source_map: dict[str, dict[str, Any]]) -> str | None:
    if not source_ids:
        return None
    parts: list[str] = []
    for sid in source_ids:
        src = source_map.get(sid)
        if not src:
            continue
        lines = [f"Источник: {src.get('name', sid)}"]
        if src.get("date"):
            lines.append(f"Дата: {src['date']}")
        if src.get("reference"):
            lines.append(f"Ссылка/место: {src['reference']}")
        lines.append(f"Подтверждение: {src.get('confidence', 'unknown')}")
        if src.get("note"):
            lines.append(str(src["note"]))
        parts.append("\n".join(lines))
    return "\n\n".join(parts) if parts else None


def _cell(
    doc: Document,
    cell,
    text: Any,
    *,
    bold: bool = False,
    size: float = 7.0,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    source_ids: list[str] | None = None,
    source_map: dict[str, dict[str, Any]] | None = None,
    comments: bool = False,
    extra_comment: str | None = None,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("—" if text in (None, "") else str(text))
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.bold = bold
    if comments:
        ctext = _comment_text(source_ids, source_map or {})
        if extra_comment:
            ctext = f"{ctext}\n\n{extra_comment}" if ctext else extra_comment
        if ctext:
            doc.add_comment(r, text=ctext, author="OSP Skill", initials="OSP")


def _header_row(doc, row, labels: list[str], size: float = 6.2) -> None:
    for i, label in enumerate(labels):
        _shade(row.cells[i], HEADER)
        _cell(doc, row.cells[i], label, bold=True, size=size, align=WD_ALIGN_PARAGRAPH.CENTER)


def _section_head(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARKBLUE)


def _title(doc: Document, number: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{number}. {text}")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARKBLUE)


def _setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Mm(297)
    sec.page_height = Mm(210)
    sec.top_margin = Mm(8)
    sec.bottom_margin = Mm(8)
    sec.left_margin = Mm(8)
    sec.right_margin = Mm(8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(7)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def _money(value: Any) -> str:
    if value in (None, ""):
        return "—"
    return f"{float(value):,.0f}".replace(",", " ")


def render_report(data: dict[str, Any], out_path: str | Path, *, include_comments: bool = True) -> Path:
    out_path = Path(out_path)
    data = inherit_parent_periods(data)
    doc = Document()
    _setup(doc)
    source_map = {s["id"]: s for s in data.get("sources", []) if s.get("id")}
    report = data["report"]
    field_sources = report.get("field_sources") or {}
    progress = calculate_progress(data)

    _title(doc, 1, "Резюме по проекту")
    passport = [
        ("Заказчик", report.get("customer"), "customer"),
        ("Проект", report.get("project_name"), "project_name"),
        ("Руководитель проекта от Заказчика", report.get("customer_pm"), "customer_pm"),
        ("Основание", report.get("contract_basis"), "contract_basis"),
        ("Дополнительное соглашение", report.get("additional_agreement"), "additional_agreement"),
        ("Номер отчета", report.get("report_number"), "report_number"),
        ("Руководитель проекта от Исполнителя", report.get("contractor_pm"), "contractor_pm"),
        ("Отчетный период", report.get("reporting_period"), "reporting_period"),
    ]
    t = doc.add_table(rows=len(passport), cols=2)
    _compact(t); _widths(t, [61, 215])
    for i, (label, value, key) in enumerate(passport):
        _shade(t.cell(i, 0), LIGHT)
        _cell(doc, t.cell(i, 0), label, bold=True, size=6.8)
        _cell(doc, t.cell(i, 1), value, size=6.8, source_ids=field_sources.get(key), source_map=source_map, comments=include_comments)
        if value in (None, ""):
            _shade(t.cell(i, 1), YELLOW)

    _section_head(doc, "СВОДНЫЕ ДАННЫЕ ПО ПРОЕКТУ")
    s = doc.add_table(rows=4, cols=2)
    _compact(s); _widths(s, [61, 215])
    summary_rows = [
        ("Плановый период этапа", report.get("planned_stage_period")),
        ("Фактический период / состояние", report.get("actual_stage_period")),
        ("Этап проекта", report.get("stage_name")),
        ("Общий статус проекта", "\n".join(f"• {x}" for x in report.get("overall_status_points", []))),
    ]
    for i, (label, value) in enumerate(summary_rows):
        _shade(s.cell(i, 0), LIGHT)
        _cell(doc, s.cell(i, 0), label, bold=True, size=6.6)
        _cell(doc, s.cell(i, 1), value, size=6.45)

    _section_head(doc, "ПЛАН И СТАТУС ЭТАПОВ")
    st = doc.add_table(rows=1 + len(data.get("stages", [])), cols=7)
    _compact(st); _widths(st, [72, 23, 27, 28, 22, 31, 73])
    _header_row(doc, st.rows[0], ["Этап", "Трудоемкость", "Бюджет", "Статус", "Коэф.", "Освоено", "Комментарий"], 5.9)
    for i, item in enumerate(data.get("stages", []), 1):
        budget = item.get("budget")
        coef = coefficient(item.get("status")) if budget is not None else None
        earned = float(budget) * coef if budget is not None else None
        vals = [
            item.get("title"), item.get("duration"), _money(budget), item.get("status"),
            f"{coef * 100:.0f}%" if coef is not None else "—", _money(earned), item.get("comment"),
        ]
        for j, val in enumerate(vals):
            _cell(doc, st.cell(i, j), val, size=5.7, source_ids=item.get("sources"), source_map=source_map, comments=include_comments)
        if item.get("requires_confirmation"):
            _shade(st.cell(i, 3), YELLOW)

    _section_head(doc, "ПРОГРЕСС И БЮДЖЕТ")
    bt = doc.add_table(rows=2, cols=5)
    _compact(bt); _widths(bt, [56, 56, 56, 56, 52])
    _header_row(doc, bt.rows[0], ["Утвержденный бюджет", "Расчетное освоение", "Остаток", "Прогресс", "База расчета"], 6.1)
    vals = [
        _money(progress.approved_budget),
        _money(progress.earned),
        _money(progress.approved_budget - progress.earned),
        f"{progress.percent:.1f}%",
        "Задачи" if progress.basis == "tasks" else "Этапы",
    ]
    for j, val in enumerate(vals):
        _cell(doc, bt.cell(1, j), val, bold=(j == 3), size=6.8, align=WD_ALIGN_PARAGRAPH.CENTER)

    _section_head(doc, "ПЛАН / ФАКТ ОПЛАТ")
    payments = data.get("payments", [])
    pt = doc.add_table(rows=1 + len(payments), cols=7)
    _compact(pt); _widths(pt, [45, 28, 68, 31, 38, 29, 37])
    _header_row(doc, pt.rows[0], ["Платеж", "Сумма", "План / основание по ДС", "План. дата", "Факт", "Дата факта", "Статус"], 5.6)
    for i, pmt in enumerate(payments, 1):
        vals = [
            pmt.get("name"), _money(pmt.get("amount")), pmt.get("plan_basis"), pmt.get("planned_date"),
            pmt.get("actual_text"), pmt.get("actual_date"), pmt.get("status"),
        ]
        for j, val in enumerate(vals):
            _cell(doc, pt.cell(i, j), val, size=5.5, source_ids=pmt.get("sources"), source_map=source_map, comments=include_comments)
        if pmt.get("requires_confirmation"):
            _shade(pt.cell(i, 6), YELLOW)

    doc.add_page_break()
    _title(doc, 2, "Данные по задачам")
    tasks = page2_contractual_rows(data)
    tt = doc.add_table(rows=1 + len(tasks), cols=10)
    _compact(tt); _widths(tt, [10, 68, 27, 26, 25, 27, 25, 27, 27, 47])
    _header_row(doc, tt.rows[0], ["№", "Задача / результат", "Ответственный", "Статус", "Бюджет", "Освоено", "План нач.", "План зав.", "Факт", "Комментарий"], 5.5)
    for i, item in enumerate(tasks, 1):
        budget = item.get("budget")
        earned = float(budget) * coefficient(item.get("status")) if budget is not None else None
        task_text = item.get("title")
        if item.get("result"):
            task_text = f"{task_text}\nРезультат: {item.get('result')}"
        vals = [
            item.get("id"), task_text, item.get("owner"), item.get("status"), _money(budget), _money(earned),
            item.get("planned_start"), item.get("planned_end"), item.get("actual_end"), item.get("comment"),
        ]
        for j, val in enumerate(vals):
            inherited_note = None
            if j in {6, 7} and item.get("date_basis") == "parent_stage_period":
                inherited_note = (
                    "Плановый период унаследован от договорного родительского этапа "
                    f"{item.get('contract_parent_id') or item.get('parent_id')}."
                )
            _cell(
                doc, tt.cell(i, j), val, size=5.25,
                source_ids=item.get("sources"), source_map=source_map,
                comments=include_comments, extra_comment=inherited_note,
            )
        if item.get("requires_confirmation"):
            _shade(tt.cell(i, 8), YELLOW)

    doc.add_page_break()
    _title(doc, 3, "Риски проекта")
    risks = data.get("risks", [])
    rt = doc.add_table(rows=1 + len(risks), cols=9)
    _compact(rt); _widths(rt, [18, 17, 44, 22, 21, 50, 51, 49, 27])
    _header_row(doc, rt.rows[0], ["Дата", "Тип", "Отклонение / риск", "Вероятность / факт", "Влияние", "Влияние на проект", "Корректирующие мероприятия", "Результат / решение", "Статус"], 5.15)
    for i, risk in enumerate(risks, 1):
        vals = [
            risk.get("date"), risk.get("type"), risk.get("title"), risk.get("probability_or_fact"),
            risk.get("impact_degree"), risk.get("impact"), risk.get("actions"), risk.get("result"), risk.get("status"),
        ]
        for j, val in enumerate(vals):
            _cell(doc, rt.cell(i, j), val, size=4.95, source_ids=risk.get("sources"), source_map=source_map, comments=include_comments)
        status = str(risk.get("status") or "").lower()
        _shade(rt.cell(i, 8), GREEN if "закры" in status else RED if "откры" in status else YELLOW)

    doc.add_page_break()
    _title(doc, 4, "Ключевые вопросы и проблемы")
    health = data.get("health_check", [])
    ht = doc.add_table(rows=1 + len(health), cols=4)
    _compact(ht); _widths(ht, [12, 82, 34, 148])
    _header_row(doc, ht.rows[0], ["№", "Вопрос", "Ответ", "Комментарий / принятое решение"], 6.0)
    for i, item in enumerate(health, 1):
        vals = [item.get("number"), item.get("question"), item.get("answer"), item.get("comment")]
        for j, val in enumerate(vals):
            _cell(doc, ht.cell(i, j), val, size=5.8, source_ids=item.get("sources"), source_map=source_map, comments=include_comments)

    _section_head(doc, "ОТКРЫТЫЕ ВОПРОСЫ / ДЕЙСТВИЯ")
    opens = data.get("open_items", [])
    ot = doc.add_table(rows=1 + len(opens), cols=5)
    _compact(ot); _widths(ot, [69, 101, 39, 31, 36])
    _header_row(doc, ot.rows[0], ["Открытый вопрос / действие", "Необходимое действие / решение", "Ответственный", "Срок", "Статус / результат"], 5.8)
    for i, item in enumerate(opens, 1):
        vals = [item.get("item"), item.get("required_action"), item.get("owner"), item.get("due_date"), item.get("status")]
        for j, val in enumerate(vals):
            _cell(doc, ot.cell(i, j), val, size=5.6, source_ids=item.get("sources"), source_map=source_map, comments=include_comments)

    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("ОСП сформирован по проектным источникам. Подробные ссылки доступны во внутренней версии.")
        r.font.name = "Arial"
        r.font.size = Pt(5.8)
        r.font.color.rgb = RGBColor(100, 100, 100)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
